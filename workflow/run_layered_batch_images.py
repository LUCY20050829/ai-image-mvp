from __future__ import annotations

import argparse
import base64
import concurrent.futures
import http.client
import json
import mimetypes
import os
import re
import ssl
import time
import uuid
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
PROJECT = ROOT / "ai-book-project"
APP_ROOT = ROOT / "ai自动化生图-0903"
PROMPT_ASSETS = APP_ROOT / "prompt_assets"
APP_ASSETS = APP_ROOT / "assets"
FACE_LIBRARY = APP_ASSETS / "faces"
SRC = PROJECT / "source_files"
OUT = PROJECT / "generated_images" / "批量生成图"
BG_DIR = OUT / "背景图"
PERSON_DIR = OUT / "人物图"
COMP_DIR = OUT / "合成图"
FACE_DIR = OUT / "人物脸图参考"
PROMPT_DIR = OUT / "提示词"
STATE_JSON = OUT / "batch_state.json"
PROMPT_JSON = PROMPT_DIR / "layered_prompts.json"
RESULT_JSON = OUT / "batch_results.json"
ERROR_JSON = OUT / "batch_errors.json"
OUT_XLSX = SRC / "批量生成图.xlsx"

BASE_URL = os.getenv("TAL_MLOPS_BASE_URL", "http://ai-service.tal.com")
CHAT_MODEL = "deepseek-v4-flash"
IMAGE_MODEL = "gpt-image-2"

HEADERS_OUT = [
    "序号",
    "诗人",
    "诗名",
    "诗歌内容",
    "背景图提示词",
    "背景图路径",
    "人物图提示词",
    "人物脸图参考",
    "人物图路径",
    "合成提示词",
    "合成图路径",
    "是否需人物",
    "备注",
]

SYSTEM_PROMPT = """
你是 AI 交互图书项目的分层图像提示词导演，专门为 gpt-image-2 编写三类提示词：背景图、人物图、合成图。

只输出 JSON 数组，不输出解释、Markdown 或多余文字。
数组中每个元素格式：
{"source_index":1,"poet":"","title":"","background_prompt":"","person_prompt":"","composite_instruction":""}

总体硬规则：
1. 固定画风：3D中国卡通古风电影风，横构图16:9，色彩明丽、干净，饱和度适中偏高，对比度略高。
2. 不要描写人物神情，不出现“神情”二字；不要出现对“身形”的描写，不出现“身形”二字。
3. 诗词文字区字号约 22 号；如果长诗，只放最出名、最能代表主画面的几句。

背景图提示词规则：
1. 背景图只生成环境、景物、诗词文字区、必要的背景人物；不要生成最终诗人本人，不要出现诗人的脸。
2. 背景图必须给诗人预留位置，并明确 center_x、foot_y、height_ratio、朝向、落脚点或坐姿支撑。
3. 背景必须有可供人物融合的空间逻辑，比如地面、船舷、栏杆、山石、桌案、窗边、院墙等。
4. 主画面服从输入的主视觉和背景图生成说明，不要堆满所有意象。
5. 想象、怀古、历史类诗歌必须分现实层和想象/历史层；诗人位置只在现实层预留。
6. 诗词文字区优先自然嵌入天空、水面、雾气、月光、墙面、山岚、竹影等留白；必要时才用柔和白色或米白渐变。
7. 除诗词区外，不得出现其他文字、伪文字、水印、招牌、门联。

人物图提示词规则：
1. 如果 person_should_appear 为 false，person_prompt 写空字符串，composite_instruction 写“诗人不出现，合成图直接采用背景图”，不要编造人物。
2. 如果 person_should_appear 为 true，人物图基于输入诗人脸图图生图，输出独立诗人形象，用于后期合成。
3. 人物图提示词必须逐行写：
保持原图脸部特征不改变！
不要改变原图比例！！！
4. 只改变服饰、动作、朝向、手持道具；不重新设计脸。
5. 服饰根据朝代和诗歌场景改为合适常服；不要官服，不要绿色或青绿色服装。
6. 人物动作、朝向、大小必须匹配输入锚点，方便放入背景预留位置。
7. 人物图最好是完整人物或需要的半身人物，背景尽量简单干净，方便后期抠图或图生图合成。

合成图提示词规则：
1. 合成图以第一张输入图为背景图，第二张输入图为诗人人物图。
2. 把人物放到背景中预留的锚点位置，严格参考 center_x、foot_y、height_ratio、朝向和遮挡关系。
3. 保持人物脸部特征不改变，不改变原图比例；只调整整体缩放、边缘融合、光影方向和遮挡。
4. 人物脚底/坐姿底部必须与地面、船、山石、桌案或窗边自然接触，不要悬浮。
5. 诗词文字保持清楚可读，约 22 号，不要被人物遮挡。
6. 不要额外生成第二个诗人，不要生成现代元素或额外文字。

每条 prompt 必须分段，段名固定：
【任务目标】
【画风与构图】
【背景/人物/合成要求】
【诗词文字区】
【参考图使用】
【禁忌】
【诗歌原文】
""".strip()


def compact(value) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip())


def norm(value) -> str:
    text = compact(value)
    text = text.replace("九月节日忆山东兄弟", "九月九日忆山东兄弟")
    return re.sub(r"[《》\s・·/（）()【】\[\]「」『』,，.。:：;；!！?？]", "", text)


def safe_name(value: str) -> str:
    text = re.sub(r"[《》“”\"'`]+", "", compact(value))
    text = re.sub(r"[\\/:*?<>|]+", "_", text)
    text = re.sub(r"\s+", "", text)
    return text[:80].strip("_") or "item"


def ensure_dirs() -> None:
    for d in [OUT, BG_DIR, PERSON_DIR, COMP_DIR, FACE_DIR, PROMPT_DIR]:
        d.mkdir(parents=True, exist_ok=True)


def xlsx_by_name(parts: list[str]) -> Path:
    matches = [p for p in SRC.glob("*.xlsx") if all(part in p.name for part in parts) and not p.name.startswith("~$")]
    if not matches:
        raise FileNotFoundError(parts)
    return matches[0]


def read_sheet(path: Path, sheet_name: str | None = None) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    rows = []
    for r in range(2, ws.max_row + 1):
        item = {name: ws.cell(r, c).value for name, c in headers.items()}
        if any(item.values()):
            item["_excel_row"] = r
            rows.append(item)
    wb.close()
    return rows


def credentials() -> tuple[str, str]:
    app_id = os.getenv("TAL_MLOPS_APP_ID")
    app_key = os.getenv("TAL_MLOPS_APP_KEY")
    if not app_id or not app_key:
        raise RuntimeError("missing TAL_MLOPS_APP_ID/TAL_MLOPS_APP_KEY")
    return app_id, app_key


def post_json(path: str, payload: dict, headers: dict, timeout: int = 900) -> dict:
    parsed = urlparse(BASE_URL)
    cls = http.client.HTTPSConnection if parsed.scheme == "https" else http.client.HTTPConnection
    kwargs = {"timeout": timeout}
    if parsed.scheme == "https":
        kwargs["context"] = ssl.create_default_context()
    conn = cls(parsed.netloc or parsed.path, **kwargs)
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    conn.request("POST", path, body=body, headers={"Accept": "application/json", "Content-Type": "application/json", **headers})
    resp = conn.getresponse()
    data = resp.read()
    conn.close()
    if resp.status < 200 or resp.status >= 300:
        raise RuntimeError(f"HTTP {resp.status}: {data.decode('utf-8', errors='replace')}")
    return json.loads(data.decode("utf-8"))


def multipart_body(fields: dict, files: list[tuple[str, Path]]) -> tuple[bytes, str]:
    boundary = f"----ai-book-layer-batch-{uuid.uuid4().hex}"
    chunks: list[bytes] = []
    for name, value in fields.items():
        chunks.append(f"--{boundary}\r\n".encode())
        chunks.append(f'Content-Disposition: form-data; name="{name}"\r\n\r\n'.encode())
        chunks.append(str(value).encode("utf-8"))
        chunks.append(b"\r\n")
    for name, path in files:
        ctype = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
        chunks.append(f"--{boundary}\r\n".encode())
        chunks.append((f'Content-Disposition: form-data; name="{name}"; filename="{path.name}"\r\nContent-Type: {ctype}\r\n\r\n').encode("utf-8"))
        chunks.append(path.read_bytes())
        chunks.append(b"\r\n")
    chunks.append(f"--{boundary}--\r\n".encode())
    return b"".join(chunks), boundary


def post_multipart(path: str, fields: dict, files: list[tuple[str, Path]]) -> dict:
    app_id, app_key = credentials()
    body, boundary = multipart_body(fields, files)
    parsed = urlparse(BASE_URL)
    cls = http.client.HTTPSConnection if parsed.scheme == "https" else http.client.HTTPConnection
    conn = cls(parsed.netloc or parsed.path, timeout=900)
    conn.request(
        "POST",
        path,
        body=body,
        headers={
            "Accept": "application/json",
            "api-key": f"{app_id}:{app_key}",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
    )
    resp = conn.getresponse()
    data = resp.read()
    conn.close()
    if resp.status < 200 or resp.status >= 300:
        raise RuntimeError(f"HTTP {resp.status}: {data.decode('utf-8', errors='replace')}")
    return json.loads(data.decode("utf-8"))


def save_image_response(resp: dict, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    first = (resp.get("data") or [{}])[0]
    if first.get("b64_json"):
        output.write_bytes(base64.b64decode(first["b64_json"]))
        return
    if first.get("url"):
        req = Request(first["url"], headers={"User-Agent": "ai-book-layered-batch/1.0"})
        with urlopen(req, timeout=900) as r:
            output.write_bytes(r.read())
        return
    raise RuntimeError(json.dumps(resp, ensure_ascii=False))


def find_person_book() -> Path:
    candidates = list(Path("C:/Users/427483/Desktop").rglob("人物形象及人物形象资料.xlsx"))
    if not candidates:
        raise FileNotFoundError("人物形象及人物形象资料.xlsx")
    return candidates[0]


def xiaozhong_rows(person_book: Path) -> dict[str, tuple[int, dict]]:
    wb = load_workbook(person_book, read_only=True, data_only=True)
    ws = wb["小中小高"]
    headers = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    out = {}
    for r in range(2, ws.max_row + 1):
        poet = compact(ws.cell(r, headers["人物"]).value)
        if poet:
            out[poet] = (r, {name: ws.cell(r, c).value for name, c in headers.items()})
    wb.close()
    return out


def extract_face_from_book(person_book: Path, poet: str, row_number: int) -> Path | None:
    out = FACE_DIR / f"{safe_name(poet)}_face_from_xiaozhong_row{row_number}.png"
    if out.exists() and out.stat().st_size > 10000:
        return out
    wb = load_workbook(person_book)
    ws = wb["小中小高"]
    try:
        for img in ws._images:
            anchor = img.anchor._from
            if anchor.row + 1 == row_number and anchor.col + 1 == 3:
                out.write_bytes(img._data())
                return out
    finally:
        wb.close()
    return None


def preferred_face(person_book: Path, poet: str, row_number: int) -> Path | None:
    # Prefer explicitly approved/generated reusable face refs.
    face_dir = FACE_LIBRARY
    if poet == "李白":
        libai = face_dir / "libai_face_from_xinglunan_bg4_2.png"
        if libai.exists() and libai.stat().st_size > 10000:
            return libai
    patterns = [f"{poet}*face*.png", f"{poet}*脸*.png", f"{poet}_face.*"]
    for pattern in patterns:
        matches = sorted(face_dir.glob(pattern))
        for m in matches:
            if m.exists() and m.stat().st_size > 10000:
                return m
    return extract_face_from_book(person_book, poet, row_number)


def load_source_rows() -> list[dict]:
    valid_path = xlsx_by_name(["批量生成背景图4", "有效诗歌清单"])
    anchor_path = xlsx_by_name(["人物锚点前置判断"])
    valid = read_sheet(valid_path, "有效诗歌清单")
    anchors = read_sheet(anchor_path, "人物锚点前置判断")
    by_key = {(compact(r.get("诗人")), norm(r.get("诗名"))): r for r in anchors}
    rows = []
    for v in valid:
        poet = compact(v.get("诗人"))
        title = compact(v.get("诗名"))
        a = by_key.get((poet, norm(title)))
        if not a:
            # Fallback for small title normalizations.
            for (p, t), row in by_key.items():
                if p == poet and (norm(title) in t or t in norm(title)):
                    a = row
                    break
        if not a:
            rows.append({"序号": v.get("序号"), "诗人": poet, "诗名": title, "_missing_anchor": True})
            continue
        merged = dict(a)
        merged["序号"] = v.get("序号")
        rows.append(merged)
    return rows


def person_should_appear(row: dict) -> bool:
    value = row.get("诗人是否出现")
    if isinstance(value, bool):
        return value
    return compact(value).lower() in {"true", "yes", "1", "是", "需要", "出现"}


def prepare() -> dict:
    ensure_dirs()
    person_book = find_person_book()
    people = xiaozhong_rows(person_book)
    rows = load_source_rows()
    missing_anchor = [r for r in rows if r.get("_missing_anchor")]
    missing_face = []
    prepared = []
    for row in rows:
        poet = compact(row.get("诗人"))
        title = compact(row.get("诗名"))
        face = ""
        person_info = {}
        if person_should_appear(row):
            person_row = people.get(poet)
            if not person_row:
                missing_face.append({"poet": poet, "title": title, "reason": "missing poet in 小中小高"})
            else:
                rn, person_info = person_row
                face_path = preferred_face(person_book, poet, rn)
                if not face_path:
                    missing_face.append({"poet": poet, "title": title, "reason": "missing face image"})
                else:
                    face = str(face_path)
        prepared.append({"row": row, "face": face, "person_info": person_info})
    state = {
        "rows": len(rows),
        "need_person": sum(1 for r in rows if person_should_appear(r)),
        "no_person": sum(1 for r in rows if not person_should_appear(r)),
        "missing_anchor": missing_anchor,
        "missing_face": missing_face,
        "out_dir": str(OUT),
    }
    STATE_JSON.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    return state


def source_payload(item: dict) -> dict:
    row = item["row"]
    return {
        "source_index": row.get("序号"),
        "诗人": row.get("诗人"),
        "朝代": "唐" if row.get("诗人") not in {"陶渊明"} else "东晋",
        "诗名": row.get("诗名"),
        "诗歌原文": row.get("诗歌内容"),
        "诗歌释义": row.get("诗歌释义"),
        "person_should_appear": person_should_appear(row),
        "场景类型": row.get("场景类型"),
        "现实层": row.get("现实层"),
        "想象/历史层": row.get("想象/历史层"),
        "主视觉": row.get("主视觉"),
        "背景图生成说明": row.get("背景图生成说明"),
        "人物姿态": row.get("人物姿态"),
        "人物朝向": row.get("人物朝向"),
        "人物动作": row.get("人物动作"),
        "center_x": row.get("center_x"),
        "foot_y": row.get("foot_y"),
        "height_ratio": row.get("height_ratio"),
        "pose": row.get("pose"),
        "facing": row.get("facing"),
        "遮挡关系": row.get("遮挡关系"),
        "阴影方向": row.get("阴影方向"),
        "融合注意": row.get("融合注意"),
        "诗词区位置": row.get("诗词区位置"),
        "诗词区方式": row.get("诗词区方式"),
        "字号": 22,
        "建筑类别": row.get("建筑类别"),
        "建筑形制注意": row.get("建筑形制注意"),
        "花卉": row.get("花卉"),
        "花卉形态注意": row.get("花卉形态注意"),
        "避免误画": row.get("避免误画"),
        "人物图参考资料": item.get("person_info") or {},
    }


def parse_json_array(text: str) -> list[dict]:
    try:
        value = json.loads(text)
    except Exception:
        match = re.search(r"\[.*\]", text, re.S)
        if not match:
            raise
        value = json.loads(match.group(0))
    if not isinstance(value, list):
        raise ValueError("DeepSeek output is not JSON array")
    return value


def sanitize_prompt_item(item: dict) -> dict:
    def clean(s):
        if not isinstance(s, str):
            return s
        s = s.replace("18-22", "22").replace("20-22", "22").replace("24号", "22号").replace("24 号", "22 号")
        s = s.replace("神情", "").replace("身形", "人物")
        return s

    return {k: clean(v) for k, v in item.items()}


def prompt_key(row: dict) -> str:
    return f"{int(row.get('序号') or 0):03d}_{compact(row.get('诗人'))}_{compact(row.get('诗名'))}"


def ask_deepseek(batch: list[dict]) -> list[dict]:
    app_id, app_key = credentials()
    payload = {
        "model": CHAT_MODEL,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {
                "role": "user",
                "content": "请为以下诗歌各生成一组分层提示词，每首输出 background_prompt、person_prompt、composite_instruction：\n"
                + json.dumps([source_payload(x) for x in batch], ensure_ascii=False, indent=2),
            },
        ],
        "temperature": 0.2,
        "max_tokens": 12000,
    }
    resp = post_json("/openai-compatible/v1/chat/completions", payload, {"Authorization": f"Bearer {app_id}:{app_key}"}, timeout=420)
    return [sanitize_prompt_item(x) for x in parse_json_array(resp["choices"][0]["message"]["content"])]


def load_prompt_pack() -> dict:
    if PROMPT_JSON.exists():
        return json.loads(PROMPT_JSON.read_text(encoding="utf-8"))
    return {}


def save_prompt_pack(pack: dict) -> None:
    PROMPT_JSON.write_text(json.dumps(pack, ensure_ascii=False, indent=2), encoding="utf-8")


def chunked(seq: list, n: int):
    for i in range(0, len(seq), n):
        yield seq[i : i + n]


def build_items() -> list[dict]:
    person_book = find_person_book()
    people = xiaozhong_rows(person_book)
    rows = load_source_rows()
    items = []
    for row in rows:
        poet = compact(row.get("诗人"))
        face = ""
        person_info = {}
        if person_should_appear(row) and poet in people:
            rn, person_info = people[poet]
            face_path = preferred_face(person_book, poet, rn)
            if face_path:
                face = str(face_path)
        items.append({"row": row, "face": face, "person_info": person_info})
    return items


def generate_prompts(workers: int, batch_size: int) -> dict:
    ensure_dirs()
    items = build_items()
    pack = load_prompt_pack()
    todo = [x for x in items if prompt_key(x["row"]) not in pack]
    print(json.dumps({"prompt_pending": len(todo), "already": len(pack), "batch_size": batch_size, "workers": workers}, ensure_ascii=False), flush=True)
    errors = []

    def one(batch: list[dict]):
        return batch, ask_deepseek(batch)

    batches = list(chunked(todo, batch_size))
    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(one, batch): batch for batch in batches}
        for n, fut in enumerate(concurrent.futures.as_completed(futs), 1):
            batch = futs[fut]
            try:
                _, result = fut.result()
                by_source = {int(x.get("source_index") or 0): x for x in result}
                for item in batch:
                    row = item["row"]
                    key = prompt_key(row)
                    prompt = by_source.get(int(row.get("序号") or 0))
                    if not prompt:
                        # Fallback by poet/title if DeepSeek omitted source_index.
                        for p in result:
                            if compact(p.get("poet")) == compact(row.get("诗人")) and norm(p.get("title")) == norm(row.get("诗名")):
                                prompt = p
                                break
                    if prompt:
                        pack[key] = prompt
                    else:
                        errors.append({"stage": "prompt", "key": key, "error": "missing prompt in DeepSeek output"})
                save_prompt_pack(pack)
                print(f"prompt_batch {n}/{len(batches)} done total={len(pack)}", flush=True)
            except Exception as exc:
                for item in batch:
                    errors.append({"stage": "prompt", "key": prompt_key(item["row"]), "error": repr(exc)})
                ERROR_JSON.write_text(json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8")
                print("ERR " + json.dumps(errors[-1], ensure_ascii=False), flush=True)
    if errors:
        ERROR_JSON.write_text(json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8")
    return pack


def bg_path(row: dict) -> Path:
    return BG_DIR / f"{prompt_key(row)}_background.png"


def person_path(row: dict) -> Path:
    return PERSON_DIR / f"{prompt_key(row)}_person.png"


def comp_path(row: dict) -> Path:
    return COMP_DIR / f"{prompt_key(row)}_composite.png"


def image_request(prompt: str, output: Path, refs: list[Path]) -> tuple[bool, str]:
    if output.exists() and output.stat().st_size > 10000:
        return True, f"exists {output.name}"
    for attempt in range(1, 4):
        try:
            if refs:
                resp = post_multipart("/openai-compatible/v1/images/edits", {"model": IMAGE_MODEL, "prompt": prompt}, [("image[]", p) for p in refs])
            else:
                app_id, app_key = credentials()
                resp = post_json("/openai-compatible/v1/images/generations", {"model": IMAGE_MODEL, "prompt": prompt}, {"api-key": f"{app_id}:{app_key}"}, timeout=900)
            save_image_response(resp, output)
            return True, f"generated {output.name}"
        except Exception as exc:
            if attempt == 3:
                return False, f"failed {output.name}: {exc}"
            time.sleep(70 if "429" in str(exc) else 20)
    return False, f"failed {output.name}"


def generate_backgrounds(workers: int) -> list[dict]:
    items = build_items()
    pack = load_prompt_pack()
    jobs = []
    for item in items:
        row = item["row"]
        p = pack.get(prompt_key(row))
        if not p:
            continue
        out = bg_path(row)
        if out.exists() and out.stat().st_size > 10000:
            continue
        jobs.append({"row": row, "prompt": p.get("background_prompt") or "", "out": out})
    return run_image_jobs("background", jobs, workers)


def generate_people(workers: int) -> list[dict]:
    items = build_items()
    pack = load_prompt_pack()
    jobs = []
    for item in items:
        row = item["row"]
        if not person_should_appear(row):
            continue
        p = pack.get(prompt_key(row))
        if not p or not item.get("face"):
            continue
        out = person_path(row)
        if out.exists() and out.stat().st_size > 10000:
            continue
        jobs.append({"row": row, "prompt": p.get("person_prompt") or "", "out": out, "refs": [Path(item["face"])]})
    return run_image_jobs("person", jobs, workers)


def generate_composites(workers: int) -> list[dict]:
    items = build_items()
    pack = load_prompt_pack()
    jobs = []
    for item in items:
        row = item["row"]
        p = pack.get(prompt_key(row))
        if not p:
            continue
        out = comp_path(row)
        if out.exists() and out.stat().st_size > 10000:
            continue
        if not person_should_appear(row):
            bg = bg_path(row)
            if bg.exists():
                out.write_bytes(bg.read_bytes())
                continue
        bg = bg_path(row)
        person = person_path(row)
        if bg.exists() and person.exists():
            jobs.append({"row": row, "prompt": p.get("composite_instruction") or "", "out": out, "refs": [bg, person]})
    return run_image_jobs("composite", jobs, workers)


def run_image_jobs(stage: str, jobs: list[dict], workers: int) -> list[dict]:
    print(json.dumps({f"{stage}_pending": len(jobs), "workers": workers}, ensure_ascii=False), flush=True)
    results = []
    errors = []

    def one(job):
        start = time.time()
        ok, msg = image_request(job["prompt"], job["out"], job.get("refs", []))
        return {**job, "ok": ok, "message": msg, "seconds": round(time.time() - start, 1)}

    with concurrent.futures.ThreadPoolExecutor(max_workers=workers) as ex:
        futs = {ex.submit(one, job): job for job in jobs}
        for n, fut in enumerate(concurrent.futures.as_completed(futs), 1):
            try:
                res = fut.result()
                row = res["row"]
                record = {
                    "stage": stage,
                    "key": prompt_key(row),
                    "poet": row.get("诗人"),
                    "title": row.get("诗名"),
                    "output": str(res["out"]),
                    "ok": res["ok"],
                    "seconds": res["seconds"],
                    "message": res["message"],
                }
                results.append(record)
                if not res["ok"]:
                    errors.append(record)
                print(f"{stage} {n}/{len(jobs)} ok={res['ok']} {res['seconds']}s {res['out'].name}", flush=True)
            except Exception as exc:
                job = futs[fut]
                row = job["row"]
                record = {"stage": stage, "key": prompt_key(row), "ok": False, "error": repr(exc)}
                results.append(record)
                errors.append(record)
                print("ERR " + json.dumps(record, ensure_ascii=False), flush=True)
            if errors:
                ERROR_JSON.write_text(json.dumps(errors, ensure_ascii=False, indent=2), encoding="utf-8")
    existing = []
    if RESULT_JSON.exists():
        existing = json.loads(RESULT_JSON.read_text(encoding="utf-8"))
    RESULT_JSON.write_text(json.dumps(existing + results, ensure_ascii=False, indent=2), encoding="utf-8")
    return results


def make_workbook(embed: bool = False) -> None:
    ensure_dirs()
    items = build_items()
    pack = load_prompt_pack()
    wb = Workbook()
    ws = wb.active
    ws.title = "总表"
    ws.append(HEADERS_OUT)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [8, 12, 28, 48, 72, 44, 72, 44, 44, 72, 44, 12, 36]
    for c, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = width
    for idx, item in enumerate(items, 2):
        row = item["row"]
        p = pack.get(prompt_key(row), {})
        values = [
            row.get("序号"),
            row.get("诗人"),
            row.get("诗名"),
            row.get("诗歌内容"),
            p.get("background_prompt"),
            str(bg_path(row)) if bg_path(row).exists() else "",
            p.get("person_prompt"),
            item.get("face"),
            str(person_path(row)) if person_path(row).exists() else "",
            p.get("composite_instruction"),
            str(comp_path(row)) if comp_path(row).exists() else "",
            person_should_appear(row),
            "",
        ]
        for c, value in enumerate(values, 1):
            ws.cell(idx, c).value = value
        ws.row_dimensions[idx].height = 90
        if embed:
            for col, path in [(6, bg_path(row)), (9, person_path(row)), (11, comp_path(row))]:
                if path.exists() and path.stat().st_size > 10000:
                    img = XLImage(str(path))
                    img.width = 180
                    img.height = 101
                    ws.add_image(img, ws.cell(idx, col).coordinate)
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUT_XLSX)
    wb.close()


def status() -> dict:
    items = build_items()
    pack = load_prompt_pack()
    summary = {
        "rows": len(items),
        "prompts": len(pack),
        "backgrounds": len([p for p in BG_DIR.glob("*.png") if p.stat().st_size > 10000]) if BG_DIR.exists() else 0,
        "people": len([p for p in PERSON_DIR.glob("*.png") if p.stat().st_size > 10000]) if PERSON_DIR.exists() else 0,
        "composites": len([p for p in COMP_DIR.glob("*.png") if p.stat().st_size > 10000]) if COMP_DIR.exists() else 0,
        "need_person": sum(1 for x in items if person_should_appear(x["row"])),
        "out_dir": str(OUT),
        "xlsx": str(OUT_XLSX),
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=["prepare", "prompts", "backgrounds", "people", "composites", "workbook", "status", "all"], default="status")
    parser.add_argument("--prompt-workers", type=int, default=6)
    parser.add_argument("--prompt-batch-size", type=int, default=5)
    parser.add_argument("--image-workers", type=int, default=50)
    parser.add_argument("--embed", action="store_true")
    args = parser.parse_args()

    ensure_dirs()
    if args.stage == "prepare":
        print(json.dumps(prepare(), ensure_ascii=False, indent=2), flush=True)
    elif args.stage == "prompts":
        generate_prompts(args.prompt_workers, args.prompt_batch_size)
    elif args.stage == "backgrounds":
        generate_backgrounds(args.image_workers)
    elif args.stage == "people":
        generate_people(args.image_workers)
    elif args.stage == "composites":
        generate_composites(args.image_workers)
    elif args.stage == "workbook":
        make_workbook(embed=args.embed)
    elif args.stage == "all":
        print(json.dumps(prepare(), ensure_ascii=False, indent=2), flush=True)
        generate_prompts(args.prompt_workers, args.prompt_batch_size)
        generate_backgrounds(args.image_workers)
        generate_people(args.image_workers)
        generate_composites(args.image_workers)
        make_workbook(embed=args.embed)
    status()


if __name__ == "__main__":
    main()



